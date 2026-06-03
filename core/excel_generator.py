import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

class ExcelGenerator:
    def __init__(self, output_dir="exports"):
        self.output_dir = output_dir
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)

    def generate_bill_excel(self, bill, template) -> str:
        file_path = os.path.join(self.output_dir, f"{bill['bill_number']}.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = f"Invoice {bill['bill_number']}"

        # Header
        ws.merge_cells('A1:D1')
        ws['A1'] = template.company_name
        ws['A1'].font = Font(size=16, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws.append([]) # Empty row
        
        ws.append(['Bill Number:', bill['bill_number']])
        ws.append(['Date:', bill['date_time'].strftime('%Y-%m-%d %H:%M')])
        if bill['customer_name']:
            ws.append(['Customer Name:', bill['customer_name']])
        if bill['customer_phone']:
            ws.append(['Customer Phone:', bill['customer_phone']])
            
        ws.append([]) # Empty row

        # Table Header
        headers = ['Description', 'Quantity', 'Unit Price', 'Amount']
        ws.append(headers)
        for col in range(1, 5):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.font = Font(bold=True)
            
        # Items
        for item in bill['items']:
            ws.append([
                item['description'],
                item['quantity'],
                item['unit_price'],
                item['amount']
            ])
            
        ws.append([])
        
        # Totals
        ws.append(['', '', 'Subtotal:', bill['subtotal']])
        if bill['tax_amount'] > 0:
            ws.append(['', '', 'Tax:', bill['tax_amount']])
        if bill['discount_amount'] > 0:
            ws.append(['', '', 'Discount:', f"-{bill['discount_amount']}"])
            
        ws.append(['', '', 'Grand Total:', bill['grand_total']])
        
        # Bold totals
        for row in range(ws.max_row - 3, ws.max_row + 1):
            ws.cell(row=row, column=3).font = Font(bold=True)
            ws.cell(row=row, column=4).font = Font(bold=True)

        # Auto-adjust column width
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width

        wb.save(file_path)
        return file_path
